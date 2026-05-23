from io import BytesIO

def render_quote_pdf(payload: dict) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buffer = BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(50, 800, f"Quote: {payload.get('quote_no', 'N/A')}")
    pdf.setFont("Helvetica", 10)
    pdf.drawString(50, 780, f"Title: {payload.get('title', '')}")
    y = 750
    for item in payload.get("items", [])[:30]:
        pdf.drawString(50, y, f"{item.get('item_code')} {item.get('item_name')} x{item.get('quantity')}")
        y -= 16
    pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def render_quote_excel(payload: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Quote"
    ws["A1"] = "Quote No"
    ws["B1"] = payload.get("quote_no", "N/A")
    ws["A2"] = "Title"
    ws["B2"] = payload.get("title", "")
    ws["A4"] = "Code"
    ws["B4"] = "Name"
    ws["C4"] = "Qty"
    ws["D4"] = "Unit Price"
    ws["E4"] = "Line Total"
    for cell in ("A1", "A2", "A4", "B4", "C4", "D4", "E4"):
        ws[cell].font = Font(bold=True)
    row = 5
    for item in payload.get("items", []):
        ws[f"A{row}"] = item.get("item_code")
        ws[f"B{row}"] = item.get("item_name")
        ws[f"C{row}"] = float(item.get("quantity", 0))
        ws[f"D{row}"] = float(item.get("unit_price", 0))
        ws[f"E{row}"] = float(item.get("line_total", 0))
        row += 1
    for col in ("A", "B", "C", "D", "E"):
        ws[f"{col}4"].alignment = Alignment(horizontal="center")
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
