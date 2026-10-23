"""Generate project progress Word report onto Desktop."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

DESKTOP = Path(r"C:\Users\Administrator\Desktop")
OUT = DESKTOP / "振航ERP项目进度报告.docx"


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_heading_cn(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=16 if level == 1 else 13 if level == 2 else 12, bold=True)
    return p


def add_para(doc, text, bold=False, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
        for p in table.rows[0].cells[i].paragraphs:
            for run in p.runs:
                set_run_font(run, bold=True, size=10)
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            table.rows[r_idx + 1].cells[c_idx].text = str(val)
            for p in table.rows[r_idx + 1].cells[c_idx].paragraphs:
                for run in p.runs:
                    set_run_font(run, size=10)
    doc.add_paragraph()


def main() -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("振航工业数字生态系统（ERP）\n项目总体进度报告")
    set_run_font(r, size=18, bold=True)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(
        f"报告日期：{datetime.now().strftime('%Y年%m月%d日')}\n"
        "仓库：shengtaixitong（本地混合部署）"
    )
    set_run_font(r, size=10, color=RGBColor(0x66, 0x66, 0x66))

    add_heading_cn(doc, "一、总体结论", 1)
    add_para(
        doc,
        "截至本报告日，振航 ERP 已形成可本地运行的业务闭环：后台 FastAPI + PostgreSQL/Redis/MinIO，"
        "前台 Web 管理端（Next.js）具备报价、项目、主数据、供应链、财务记账、审批中心等完整菜单。"
        "综合评估：项目总体完成度约 72%～78%；其中「财务记账闭环」约 80%，「主数据与报价/项目」约 85%，"
        "「设备全生命周期（序列号/验收/安装/售后）」约 70%，「对账导出与整站回归」约 40%。",
    )
    add_para(
        doc,
        "当前状态：本地 Docker 基础设施 + 本机 API(:8000) + Web(:3000) 可启动；"
        "适合内部试用与业务验收，尚未达到「全模块全接口已签字交付」标准。",
    )

    add_heading_cn(doc, "二、系统构成", 1)
    add_table(
        doc,
        ["子系统", "技术", "入口", "状态"],
        [
            ["后端 API", "FastAPI + SQLAlchemy + Alembic", "http://localhost:8000/api/v1", "可用"],
            ["Web 管理端", "Next.js 生产 standalone", "http://localhost:3000", "可用"],
            ["基础设施", "Postgres / Redis / MinIO（Docker）", "docker compose", "可用"],
            ["移动审批", "Vite 应用", "http://localhost:5174", "代码具备，需单独启动验证"],
            ["车间终端", "Vite 应用", "http://localhost:5175", "代码具备，需单独启动验证"],
            ["启动脚本", "scripts/start-erp.ps1", "Windows 混合模式", "可用"],
        ],
    )
    add_para(
        doc,
        "数据库迁移：约 40+ 个 Alembic 版本（含记账、采购应付、设备序列号、验收、安装、售后等）。"
        "OpenAPI 约 170+ 路由。Web 管理端约 35～40 个业务页面。",
    )

    add_heading_cn(doc, "三、分模块进度明细", 1)

    add_heading_cn(doc, "3.1 平台与权限", 2)
    add_table(
        doc,
        ["能力", "完成度", "说明"],
        [
            ["登录 / JWT / RBAC", "90%", "admin 超管；试用角色脚本可用"],
            ["审批中心", "85%", "报价审批、应收优惠、采购应付已接入"],
            ["通知中心", "80%", "列表/未读；审批相关通知已用"],
            ["审计日志 / 文件中心", "75%", "功能具备，需业务侧验收"],
            ["业务设置开关", "90%", "记账/报价应收/采购应付审批开关"],
        ],
    )

    add_heading_cn(doc, "3.2 主数据与产品", 2)
    add_table(
        doc,
        ["能力", "完成度", "说明"],
        [
            ["客户 / 供应商 / 材料 / 标准", "90%", "CRUD 与分类树可用"],
            ["产品库 / 成本售价模板", "85%", "主流程可用"],
            ["机型 / BOM", "85%", "侧栏可见，业务可用"],
            ["选项组 / 设备序列号", "70%", "近期扩展，待深度验收"],
            ["验收 / 安装调试", "70%", "页面与接口已接线"],
        ],
    )

    add_heading_cn(doc, "3.3 报价与项目", 2)
    add_table(
        doc,
        ["能力", "完成度", "说明"],
        [
            ["报价中心 / 版本 / 审批 / 导出", "85%", "核心演示链路已具备"],
            ["对客报价", "90%", "含产品/电机单价、应收联动"],
            ["项目台账 / 成本核算 / 订单利润", "80%", "与实际成本、采购转入联动"],
        ],
    )

    add_heading_cn(doc, "3.4 供应链与生产", 2)
    add_table(
        doc,
        ["能力", "完成度", "说明"],
        [
            ["采购管理", "85%", "计划→下单→到货；可提交应付审批"],
            ["车间工单", "75%", "列表与状态推进可用"],
            ["现场作业", "85%", "维修/缺料/装车；维修可生成应收"],
            ["售后服务", "70%", "工单 CRUD 已接线"],
            ["车间看板", "70%", "看板接口存在，需单独端验收"],
        ],
    )

    add_heading_cn(doc, "3.5 财务记账（近期重点）", 2)
    add_table(
        doc,
        ["阶段", "内容", "状态"],
        [
            ["P0", "记账核心表、开关、收支台账", "已完成"],
            ["P1", "老板补录、客户筛选、确认/作废", "已完成"],
            ["P2", "报价→应收、折扣审批、销售折扣上限、审批中心", "已完成"],
            ["L1/L2", "客户账目明细、欠款、收款核销、手录客户/账单", "已完成"],
            ["L3", "外出维修→应收", "已完成"],
            ["P3", "采购审批→应付", "已完成"],
            ["供应商台账", "供应商侧明细/付款（页面与接口已建）", "待验收"],
            ["P4", "对账单导出、月度汇总", "未开始"],
        ],
    )
    add_para(doc, "财务中心菜单：账目明细、收支台账、供应商台账（受 accounting_enabled 开关控制）。")

    add_heading_cn(doc, "四、综合进度估算", 1)
    add_table(
        doc,
        ["维度", "权重（约）", "完成度", "加权贡献"],
        [
            ["平台与权限", "15%", "85%", "12.8%"],
            ["主数据与产品", "15%", "82%", "12.3%"],
            ["报价与项目", "20%", "85%", "17.0%"],
            ["供应链与生产", "15%", "78%", "11.7%"],
            ["财务记账闭环", "20%", "80%", "16.0%"],
            ["移动/车间/MES/AI", "10%", "45%", "4.5%"],
            ["部署交付与整站回归", "5%", "50%", "2.5%"],
        ],
    )
    add_para(
        doc,
        "加权合计约 76.8%。对外表述建议：整体约 75%，核心业务可试用，交付签字前仍需验收与补齐导出/回归。",
        bold=True,
    )

    add_heading_cn(doc, "五、已具备的业务闭环（可演示）", 1)
    add_para(doc, "1）对客报价 → 申请/确认应收 → 账目明细欠款 → 登记收款核销")
    add_para(doc, "2）外出维修登记费用 → 生成应收 → 账目明细可见 → 收款核销")
    add_para(doc, "3）采购推进 → 提交应付 → 审批中心通过 → 收支/供应商台账应付")
    add_para(doc, "4）报价审批、应收优惠审批、采购应付审批统一进入审批中心")
    add_para(doc, "5）主数据维护、项目成本、采购到货转实际成本等既有链路")

    add_heading_cn(doc, "六、主要风险与缺口", 1)
    add_table(
        doc,
        ["风险/缺口", "影响", "建议"],
        [
            ["未做整站冒烟回归", "部分页面可能存在权限或接口细节问题", "按菜单做验收清单逐页点通"],
            ["对账导出（P4）未做", "财务对账仍依赖手工导出/截图", "排期做客户/供应商对账单 Excel"],
            ["移动端/车间终端未本轮重点验证", "多端体验不确定", "单独启动 5174/5175 验收"],
            ["生产级密钥与密码策略", "安全交付项", "按 DELIVERY_ACCEPTANCE 清单整改"],
            ["Web 为 production build", "改代码后需 rebuild 才生效", "开发期可用 dev，验收前 rebuild"],
            ["功能开关关闭时菜单隐藏", "易误判为「功能没了」", "验收前确认业务设置三开关"],
        ],
    )

    add_heading_cn(doc, "七、下一步工作建议", 1)
    add_heading_cn(doc, "7.1 立即（本周）", 2)
    add_para(doc, "1. 业务验收：按财务闭环三条路径 + 采购应付路径签字确认。")
    add_para(doc, "2. 供应商台账验收：应付列表、付款核销、欠款汇总是否符合预期。")
    add_para(doc, "3. 权限抽查：boss / sales / finance 试用账号菜单与按钮是否正确。")

    add_heading_cn(doc, "7.2 短期（1～2 周）", 2)
    add_para(doc, "1. P4：客户/供应商对账单导出、月度应收应付汇总。")
    add_para(doc, "2. 整站冒烟：看板、报价、项目、采购、工单、售后、文件、审计。")
    add_para(doc, "3. 装车登记是否需要收费入账（业务确认后实现）。")

    add_heading_cn(doc, "7.3 中期", 2)
    add_para(doc, "1. 生产部署加固：密钥、Nginx、备份、metrics 隔离（见企业部署文档）。")
    add_para(doc, "2. 移动审批与车间终端正式纳入试运行。")
    add_para(doc, "3. MES/设备数据采集按现场需求渐进接入。")

    add_heading_cn(doc, "八、试用登录说明", 1)
    add_para(doc, "Web：http://localhost:3000/login")
    add_para(doc, "管理员账号：admin（密码以现场重置结果为准，勿将弱口令用于生产）。")
    add_para(doc, "业务设置建议开启：记账功能、对客报价生成应收、采购应付审批。")

    add_heading_cn(doc, "九、结论", 1)
    add_para(
        doc,
        "振航 ERP 已从「后端平台」演进为「可试用的工业业务系统」，财务应收应付主链路打通，"
        "具备内部试运行条件。距完整交付仍差：对账导出、多端验收、生产安全清单与整站回归。"
        "建议以「业务验收签字 → P4 导出 → 交付清单」顺序推进。",
    )

    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("—— 报告结束 ——")
    set_run_font(r, size=10, color=RGBColor(0x88, 0x88, 0x88))

    DESKTOP.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)
    print("size", OUT.stat().st_size)


if __name__ == "__main__":
    main()
