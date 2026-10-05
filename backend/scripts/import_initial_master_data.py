"""
Import initial master data from CSV or Excel.

Supports: customers, machines, boms, bom_items, users.
Does not modify or replace seed_demo_quote.py demo data.

Examples:
  python scripts/import_initial_master_data.py --type customers --file data/import_samples/customers.csv
  python scripts/import_initial_master_data.py --dir data/import_samples --dry-run
  python scripts/import_initial_master_data.py --file data/import_samples/master.xlsx --dry-run
"""

from __future__ import annotations

import argparse
import csv
import sys
from collections import defaultdict
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from uuid import UUID

from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker

from _script_utils import BACKEND_ROOT, database_url, parse_bool
from app.core.security import get_password_hash
from app.models.customer import Customer
from app.models.department import Department
from app.models.machine import Machine, MachineBom, MachineBomItem, MachineCategory
from app.models.permission import Permission
from app.models.role import Role
from app.models.user import User

IMPORT_ORDER = ("customers", "machines", "boms", "bom_items", "users")
FILE_MAP = {
    "customers": "customers.csv",
    "machines": "machines.csv",
    "boms": "boms.csv",
    "bom_items": "bom_items.csv",
    "users": "users.csv",
}
EXCEL_SHEETS = {
    "customers": "customers",
    "machines": "machines",
    "boms": "boms",
    "bom_items": "bom_items",
    "users": "users",
}


class DryRunContext:
    """Track entities that would be created in the same dry-run batch."""

    planned_machine_codes: set[str] = set()
    planned_bom_keys: set[tuple[str, str]] = set()


class RowError(Exception):
    pass


class ImportReport:
    def __init__(self) -> None:
        self.created = 0
        self.skipped_duplicate = 0
        self.errors: list[str] = []

    def ok(self, msg: str) -> None:
        print(f"  + {msg}")

    def skip(self, msg: str) -> None:
        self.skipped_duplicate += 1
        print(f"  ~ skip: {msg}")

    def fail(self, line_no: int, message: str) -> None:
        self.errors.append(f"line {line_no}: {message}")
        print(f"  ! ERROR line {line_no}: {message}", file=sys.stderr)


def _read_table(path: Path, sheet: str | None = None) -> list[dict[str, str]]:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RowError("openpyxl required for Excel import") from exc
        wb = load_workbook(path, read_only=True, data_only=True)
        ws_name = sheet or wb.sheetnames[0]
        if ws_name not in wb.sheetnames:
            raise RowError(f"sheet '{ws_name}' not found in {path.name}")
        ws = wb[ws_name]
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            return []
        headers = [str(h).strip() if h is not None else "" for h in header_row]
        rows: list[dict[str, str]] = []
        for raw in rows_iter:
            if raw is None or all(c is None or str(c).strip() == "" for c in raw):
                continue
            item = {
                headers[i]: ("" if raw[i] is None else str(raw[i]).strip())
                for i in range(len(headers))
                if headers[i]
            }
            rows.append(item)
        wb.close()
        return rows
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as fh:
            lines = [ln for ln in fh if not ln.lstrip().startswith("#")]
        reader = csv.DictReader(lines)
        return [{k: (v or "").strip() for k, v in row.items() if k} for row in reader]
    raise RowError(f"unsupported file type: {suffix}")


def _require(row: dict[str, str], *keys: str, line_no: int) -> dict[str, str]:
    missing = [k for k in keys if not str(row.get(k, "")).strip()]
    if missing:
        raise RowError(f"missing required column(s): {', '.join(missing)}")
    return row


def _decimal(value: str, default: str = "0") -> Decimal:
    try:
        return Decimal(str(value or default))
    except InvalidOperation as exc:
        raise RowError(f"invalid number: {value}") from exc


def _get_admin_id(db: Session) -> UUID:
    admin = db.execute(select(User).where(User.username == "admin")).scalar_one_or_none()
    if admin is None:
        raise RowError("admin user not found — run bootstrap_admin.py first")
    return admin.id


def import_customers(db: Session, rows: list[dict[str, str]], dry_run: bool, report: ImportReport) -> None:
    admin_id = _get_admin_id(db)
    for idx, raw in enumerate(rows, start=2):
        try:
            row = _require(raw, "code", "name", line_no=idx)
            code = row["code"].strip()
            existing = db.execute(select(Customer).where(Customer.code == code, Customer.is_deleted.is_(False))).scalar_one_or_none()
            if existing:
                report.skip(f"customer {code} already exists")
                continue
            if dry_run:
                report.ok(f"[dry-run] would create customer {code}")
                report.created += 1
                continue
            db.add(
                Customer(
                    code=code,
                    name=row["name"].strip(),
                    short_name=row.get("short_name") or None,
                    industry=row.get("industry") or None,
                    level=row.get("level") or None,
                    status=row.get("status") or "active",
                    phone=row.get("phone") or None,
                    email=row.get("email") or None,
                    address=row.get("address") or None,
                    remarks=row.get("remarks") or None,
                    created_by=admin_id,
                )
            )
            report.created += 1
            report.ok(f"customer {code}")
        except RowError as exc:
            report.fail(idx, str(exc))


def _ensure_category(db: Session, code: str, admin_id: UUID) -> MachineCategory:
    cat = db.execute(select(MachineCategory).where(MachineCategory.code == code)).scalar_one_or_none()
    if cat:
        return cat
    cat = MachineCategory(code=code, name=code, description="imported", created_by=admin_id)
    db.add(cat)
    db.flush()
    return cat


def import_machines(db: Session, rows: list[dict[str, str]], dry_run: bool, report: ImportReport) -> None:
    admin_id = _get_admin_id(db)
    for idx, raw in enumerate(rows, start=2):
        try:
            row = _require(raw, "code", "name", line_no=idx)
            code = row["code"].strip()
            machine_code = (row.get("machine_code") or code).strip()
            existing = db.execute(
                select(Machine).where(Machine.code == code, Machine.is_deleted.is_(False))
            ).scalar_one_or_none()
            if existing:
                report.skip(f"machine {code} already exists")
                continue
            category_id = None
            cat_code = (row.get("category_code") or "").strip()
            if cat_code and not dry_run:
                category_id = _ensure_category(db, cat_code, admin_id).id
            elif cat_code and dry_run:
                report.ok(f"[dry-run] would ensure category {cat_code}")
            if dry_run:
                DryRunContext.planned_machine_codes.add(machine_code)
                report.ok(f"[dry-run] would create machine {code}")
                report.created += 1
                continue
            db.add(
                Machine(
                    code=code,
                    machine_code=machine_code,
                    name=row["name"].strip(),
                    model_no=row.get("model_no") or None,
                    status=row.get("status") or "draft",
                    description=row.get("description") or None,
                    category_id=category_id,
                    created_by=admin_id,
                )
            )
            report.created += 1
            report.ok(f"machine {code}")
        except RowError as exc:
            report.fail(idx, str(exc))


def import_boms(db: Session, rows: list[dict[str, str]], dry_run: bool, report: ImportReport) -> None:
    admin_id = _get_admin_id(db)
    for idx, raw in enumerate(rows, start=2):
        try:
            row = _require(raw, "machine_code", "version", line_no=idx)
            machine_code = row["machine_code"].strip()
            version = row["version"].strip()
            machine = db.execute(
                select(Machine).where(Machine.machine_code == machine_code, Machine.is_deleted.is_(False))
            ).scalar_one_or_none()
            if machine is None and dry_run and machine_code in DryRunContext.planned_machine_codes:
                DryRunContext.planned_bom_keys.add((machine_code, version))
                report.ok(f"[dry-run] would create BOM {machine_code}/{version}")
                report.created += 1
                continue
            if machine is None:
                raise RowError(f"machine not found: {machine_code}")
            existing = db.execute(
                select(MachineBom).where(
                    MachineBom.machine_id == machine.id,
                    MachineBom.version == version,
                    MachineBom.is_deleted.is_(False),
                )
            ).scalar_one_or_none()
            if existing:
                report.skip(f"BOM {machine_code}/{version} already exists")
                continue
            is_active = parse_bool(row.get("is_active"), default=False)
            if dry_run:
                report.ok(f"[dry-run] would create BOM {machine_code}/{version}")
                report.created += 1
                continue
            if is_active:
                for other in db.execute(
                    select(MachineBom).where(MachineBom.machine_id == machine.id, MachineBom.is_deleted.is_(False))
                ).scalars():
                    other.is_active = False
            bom = MachineBom(
                machine_id=machine.id,
                version=version,
                is_active=is_active,
                effective_date=date.today(),
                status=row.get("status") or "draft",
                remark=row.get("remark") or None,
                created_by=admin_id,
            )
            db.add(bom)
            report.created += 1
            report.ok(f"BOM {machine_code}/{version}")
        except RowError as exc:
            report.fail(idx, str(exc))


def import_bom_items(db: Session, rows: list[dict[str, str]], dry_run: bool, report: ImportReport) -> None:
    admin_id = _get_admin_id(db)
    totals: dict[UUID, Decimal] = defaultdict(lambda: Decimal("0"))
    for idx, raw in enumerate(rows, start=2):
        try:
            row = _require(raw, "machine_code", "bom_version", "item_code", "item_name", "unit", line_no=idx)
            machine_code = row["machine_code"].strip()
            version = row["bom_version"].strip()
            item_code = row["item_code"].strip()
            machine = db.execute(
                select(Machine).where(Machine.machine_code == machine_code, Machine.is_deleted.is_(False))
            ).scalar_one_or_none()
            bom_key = (machine_code, version)
            if machine is None and dry_run and (
                machine_code in DryRunContext.planned_machine_codes
                or bom_key in DryRunContext.planned_bom_keys
            ):
                report.ok(f"[dry-run] would create item {item_code} on {machine_code}/{version}")
                report.created += 1
                continue
            if machine is None:
                raise RowError(f"machine not found: {machine_code}")
            bom = db.execute(
                select(MachineBom).where(
                    MachineBom.machine_id == machine.id,
                    MachineBom.version == version,
                    MachineBom.is_deleted.is_(False),
                )
            ).scalar_one_or_none()
            if bom is None and dry_run and bom_key in DryRunContext.planned_bom_keys:
                report.ok(f"[dry-run] would create item {item_code} on {machine_code}/{version}")
                report.created += 1
                continue
            if bom is None:
                raise RowError(f"BOM not found: {machine_code}/{version}")
            dup = db.execute(
                select(MachineBomItem).where(
                    MachineBomItem.bom_id == bom.id,
                    MachineBomItem.item_code == item_code,
                    MachineBomItem.is_deleted.is_(False),
                )
            ).scalar_one_or_none()
            if dup:
                report.skip(f"BOM item {machine_code}/{version}/{item_code} already exists")
                continue
            qty = _decimal(row.get("quantity") or "1")
            unit_cost = _decimal(row.get("unit_cost") or "0")
            line_total = (qty * unit_cost).quantize(Decimal("0.0001"))
            sort_order = int(row.get("sort_order") or "0")
            if dry_run:
                report.ok(f"[dry-run] would create item {item_code} on {machine_code}/{version}")
                report.created += 1
                totals[bom.id] += line_total
                continue
            db.add(
                MachineBomItem(
                    bom_id=bom.id,
                    item_code=item_code,
                    item_name=row["item_name"].strip(),
                    spec=row.get("spec") or None,
                    unit=row["unit"].strip(),
                    quantity=qty,
                    unit_cost=unit_cost,
                    line_total=line_total,
                    sort_order=sort_order,
                    remark=row.get("remark") or None,
                    created_by=admin_id,
                )
            )
            totals[bom.id] += line_total
            report.created += 1
            report.ok(f"item {machine_code}/{version}/{item_code}")
        except RowError as exc:
            report.fail(idx, str(exc))
    if not dry_run:
        for bom_id, total in totals.items():
            bom = db.get(MachineBom, bom_id)
            if bom:
                bom.total_cost = total


def _ensure_permission(db: Session, code: str, name: str, resource: str, action: str) -> Permission:
    perm = db.execute(select(Permission).where(Permission.code == code)).scalar_one_or_none()
    if perm:
        return perm
    perm = Permission(name=name, code=code, resource=resource, action=action, description="imported")
    db.add(perm)
    db.flush()
    return perm


def _ensure_role(db: Session, role_code: str) -> Role | None:
    if not role_code:
        return None
    role = db.execute(select(Role).where(Role.code == role_code)).scalar_one_or_none()
    return role


def import_users(db: Session, rows: list[dict[str, str]], dry_run: bool, report: ImportReport) -> None:
    dept = db.execute(select(Department).where(Department.code == "ROOT")).scalar_one_or_none()
    if dept is None:
        raise RowError("department ROOT not found — run bootstrap_admin.py first")
    for idx, raw in enumerate(rows, start=2):
        try:
            row = _require(raw, "username", "password", line_no=idx)
            username = row["username"].strip()
            existing = db.execute(select(User).where(User.username == username)).scalar_one_or_none()
            if existing:
                report.skip(f"user {username} already exists")
                continue
            role_code = (row.get("role_code") or "").strip()
            role = _ensure_role(db, role_code) if role_code else None
            if role_code and role is None:
                raise RowError(f"role not found: {role_code} — run setup_trial_users.py first")
            email = (row.get("email") or f"{username}@trial.local").strip()
            if dry_run:
                report.ok(f"[dry-run] would create user {username}" + (f" role={role_code}" if role_code else ""))
                report.created += 1
                continue
            user = User(
                username=username,
                email=email,
                hashed_password=get_password_hash(row["password"].strip()),
                full_name=row.get("full_name") or username,
                phone=row.get("phone") or None,
                is_active=True,
                is_superuser=False,
                department_id=dept.id,
            )
            if role:
                user.roles.append(role)
            db.add(user)
            report.created += 1
            report.ok(f"user {username}")
        except RowError as exc:
            report.fail(idx, str(exc))


IMPORTERS = {
    "customers": import_customers,
    "machines": import_machines,
    "boms": import_boms,
    "bom_items": import_bom_items,
    "users": import_users,
}


def run_import(db: Session, entity: str, path: Path, sheet: str | None, dry_run: bool) -> ImportReport:
    print(f"\n== Import {entity} from {path.name}" + (f" [{sheet}]" if sheet else "") + (" [DRY-RUN]" if dry_run else ""))
    rows = _read_table(path, sheet=sheet)
    if not rows:
        print("  (empty file)")
        return ImportReport()
    report = ImportReport()
    IMPORTERS[entity](db, rows, dry_run, report)
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description="Import initial master data (CSV/Excel)")
    parser.add_argument("--type", choices=[*IMPORT_ORDER, "all"], help="entity type to import")
    parser.add_argument("--file", type=Path, help="CSV or Excel file path")
    parser.add_argument("--dir", type=Path, help="directory with standard CSV filenames")
    parser.add_argument("--sheet", help="Excel sheet name (when --type set)")
    parser.add_argument("--dry-run", action="store_true", help="validate only, no DB writes")
    args = parser.parse_args()

    if not args.type and not args.dir and not args.file:
        parser.error("specify --type + --file, or --dir, or --file with all sheets")

    engine = create_engine(database_url(), pool_pre_ping=True)
    SessionLocal = sessionmaker(bind=engine, expire_on_commit=False)

    total_errors: list[str] = []
    total_created = 0
    total_skipped = 0

    DryRunContext.planned_machine_codes = set()
    DryRunContext.planned_bom_keys = set()

    with SessionLocal() as db:
        jobs: list[tuple[str, Path, str | None]] = []

        if args.dir:
            for entity in IMPORT_ORDER:
                candidate = args.dir / FILE_MAP[entity]
                if candidate.exists():
                    jobs.append((entity, candidate, None))
        elif args.file:
            path = args.file if args.file.is_absolute() else BACKEND_ROOT / args.file
            suffix = path.suffix.lower()
            if suffix in {".xlsx", ".xlsm"}:
                entity = args.type or "all"
                if entity == "all":
                    for ent, sheet in EXCEL_SHEETS.items():
                        jobs.append((ent, path, sheet))
                else:
                    jobs.append((entity, path, args.sheet or EXCEL_SHEETS[entity]))
            else:
                if not args.type:
                    parser.error("--type required for CSV files")
                jobs.append((args.type, path, None))
        elif args.type == "all":
            parser.error("--dir or --file required for --type all")

        for entity, path, sheet in jobs:
            if not path.exists():
                print(f"SKIP missing file: {path}")
                continue
            report = run_import(db, entity, path, sheet, args.dry_run)
            total_created += report.created
            total_skipped += report.skipped_duplicate
            total_errors.extend(report.errors)

        if args.dry_run:
            db.rollback()
            print("\n[DRY-RUN] no changes committed")
        else:
            if total_errors:
                db.rollback()
                print("\nROLLED BACK due to errors")
            else:
                db.commit()
                print("\nCommitted successfully")

    print(f"\nSummary: created={total_created} skipped_duplicates={total_skipped} errors={len(total_errors)}")
    if total_errors:
        print("\nError lines:")
        for err in total_errors:
            print(f"  {err}")
        sys.exit(1)


if __name__ == "__main__":
    main()
